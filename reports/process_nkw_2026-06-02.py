#!/usr/bin/env python3
"""Weekly Negative Keyword Analysis — 2026-06-02"""
import json, os
from collections import defaultdict
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

DATA_DIR = "/root/.claude/projects/-home-user-routines-ppc/c65e0a2d-325f-4d9f-a440-638dcfffbb5f/tool-results"
REPORT_DATE = "2026-06-02"
DATE_RANGE = "May 26–June 1, 2026"

# ── tab order ────────────────────────────────────────────────────────────────
TAB_ORDER = ["ASRV","Griffin","Velo","Sunrise","TriActive","LEARN","Eurolog","Spotlight",
             "CQC","EcoPeaCo CA","EcoPeaCo US","GoGather","RCM","Ripco","We Buy"]

CLIENT_DISPLAY = {
    "ASRV":"ASRV","Griffin":"Griffin Funding","Velo":"Velo Property Solutions",
    "Sunrise":"Sunrise Yacht Products","TriActive":"TriActiveUSA","LEARN":"LEARN Behavioral",
    "Eurolog":"Eurolog Packing Group","Spotlight":"Spotlight.Vegas",
    "CQC":"Custom Quality Covers","EcoPeaCo CA":"EcoPeaCo Canada",
    "EcoPeaCo US":"EcoPeaCo USA","GoGather":"GoGather",
    "RCM":"Renewable Credit Management","Ripco":"Ripco Products","We Buy":"We Buy Exotics"
}

# ── disk files for large accounts ────────────────────────────────────────────
DISK_FILES = {
    "LEARN":        "mcp-Windsor-ai-get_data-1780395066935.txt",
    "Griffin":      "mcp-Windsor-ai-get_data-1780395068242.txt",
    "EcoPeaCo CA":  "mcp-Windsor-ai-get_data-1780395073815.txt",
    "EcoPeaCo US":  "mcp-Windsor-ai-get_data-1780395075746.txt",
    "GoGather":     "toolu_01HtARp9nWJhXXnk5ns4S59z.txt",
    "RCM":          "mcp-Windsor-ai-get_data-1780395077230.txt",
    "We Buy":       "toolu_01SDbKgsjkxGLGwtoWnYKmoX.txt",
}

# ── pre-categorized data for small inline accounts ───────────────────────────
# total_analyzed = unique terms passing threshold (clicks>=1 OR spend>=1.0)
PRECATEGORIZED = {
    "ASRV": {
        "total_analyzed": 0,
        "note": "No data returned — Windsor.ai session expired for account 160-720-1558",
        "flagged": []
    },
    "Eurolog": {
        "total_analyzed": 4,
        "flagged": [
            {"category":"ELIMINATE","search_term":"how do you insulate a shipping container",
             "campaigns":"BFD_NB_AllCampaigns","impressions":2,"clicks":1,"spend":5.9975,
             "conversions":0,"conversion_value":0,
             "reason":"DIY 'how to' informational query — B2C consumer installing insulation, not a B2B thermal packaging buyer"},
            {"category":"CONSIDER","search_term":"container blanket",
             "campaigns":"BFD_NB_AllCampaigns","impressions":2,"clicks":1,"spend":4.4703,
             "conversions":0,"conversion_value":0,
             "reason":"Ambiguous — could be B2B thermal blanket for shipping containers or consumer DIY container insulation product"},
        ]
    },
    "Sunrise": {
        "total_analyzed": 9,
        "flagged": [
            {"category":"CONSIDER","search_term":"hammock net",
             "campaigns":"multihull nets sales","impressions":1,"clicks":1,"spend":0.48,
             "conversions":0,"conversion_value":0,
             "reason":"Generic term may attract indoor/backyard hammock buyers outside Sunrise's marine netting market"},
            {"category":"CONSIDER","search_term":"boat dock netting",
             "campaigns":"multihull nets sales","impressions":1,"clicks":1,"spend":0.61,
             "conversions":0,"conversion_value":0,
             "reason":"Could attract fish trap or safety dock netting buyers rather than Sunrise's architectural dock net product"},
        ]
    },
    "TriActive": {
        "total_analyzed": 21,
        "flagged": [
            {"category":"CONSIDER","search_term":"outdoor gym equipment",
             "campaigns":"SMM_NBS_Primary_KWs","impressions":64,"clicks":8,"spend":11.1067,
             "conversions":0,"conversion_value":0,
             "reason":"Most-clicked non-converting term; generic modifier reaches residential consumers alongside target B2G buyers"},
            {"category":"CONSIDER","search_term":"outdoor gym",
             "campaigns":"SMM_NBS_Primary_KWs","impressions":35,"clicks":2,"spend":3.8281,
             "conversions":0,"conversion_value":0,
             "reason":"Very broad; captures residential/B2C consumers alongside target institutional buyers"},
            {"category":"CONSIDER","search_term":"outdoor elliptical",
             "campaigns":"SMM_NBS_Primary_KWs","impressions":1,"clicks":1,"spend":5.17,
             "conversions":0,"conversion_value":0,
             "reason":"Single-unit consumer product term; likely residential buyer not government procurement"},
            {"category":"CONSIDER","search_term":"outdoor gym setup",
             "campaigns":"SMM_NBS_Primary_KWs","impressions":5,"clicks":1,"spend":2.01,
             "conversions":0,"conversion_value":0,
             "reason":"'Setup' implies individual residential installation rather than institutional procurement"},
            {"category":"CONSIDER","search_term":"weather proof gym equipment",
             "campaigns":"SMM_NBS_Primary_KWs","impressions":1,"clicks":1,"spend":1.93,
             "conversions":0,"conversion_value":0,
             "reason":"Generic; attracts residential and institutional buyers without B2G signals"},
            {"category":"CONSIDER","search_term":"weatherproof outdoor gym equipment",
             "campaigns":"SMM_NBS_Primary_KWs","impressions":11,"clicks":3,"spend":4.24,
             "conversions":0,"conversion_value":0,
             "reason":"No conversions after 3 clicks; mixed B2C/B2G audience"},
            {"category":"CONSIDER","search_term":"outdoor exercise equipment suppliers",
             "campaigns":"SMM_NBS_Primary_KWs","impressions":1,"clicks":1,"spend":1.28,
             "conversions":0,"conversion_value":0,
             "reason":"'Suppliers' modifier may indicate competitor research or wholesale sourcing, not institutional buyer"},
            {"category":"CONSIDER","search_term":"weatherproof outdoor gym",
             "campaigns":"SMM_NBS_Primary_KWs","impressions":11,"clicks":1,"spend":1.51,
             "conversions":0,"conversion_value":0,
             "reason":"Mixed residential/commercial audience; no B2G qualification signal"},
        ]
    },
    "Velo": {
        "total_analyzed": 8,
        "flagged": [
            {"category":"CONSIDER","search_term":"post construction cleaning",
             "campaigns":"BFD_Velo_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":68.92,
             "conversions":0,"conversion_value":0,
             "reason":"Highest CPC in account this week; exceeds $40-50 target CPA; outside core segments (office/medical/warehouse)"},
            {"category":"CONSIDER","search_term":"office cleaning services",
             "campaigns":"BFD_Velo_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":38.45,
             "conversions":0,"conversion_value":0,
             "reason":"High CPC near target CPA with zero conversion; verify ad-to-landing-page alignment"},
            {"category":"CONSIDER","search_term":"cleaning companies phoenix",
             "campaigns":"BFD_Velo_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":14.64,
             "conversions":0,"conversion_value":0,
             "reason":"Too generic; may attract residential cleaning seekers, not B2B commercial prospects"},
            {"category":"CONSIDER","search_term":"construction cleaning contractors",
             "campaigns":"BFD_Velo_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":11.58,
             "conversions":0,"conversion_value":0,
             "reason":"Construction cleaning is outside Velo's core segments; review whether to exclude"},
            {"category":"CONSIDER","search_term":"commercial restaurant cleaning",
             "campaigns":"BFD_Velo_Google_NB_Search_Pros_T3","impressions":1,"clicks":1,"spend":6.66,
             "conversions":0,"conversion_value":0,
             "reason":"Restaurant cleaning is not in Velo's core segments (office/medical/warehouse/industrial)"},
        ]
    },
    "CQC": {
        "total_analyzed": 36,
        "flagged": [
            {"category":"COMPETITOR","search_term":"modz seat covers",
             "campaigns":"CQC_NBS","impressions":11,"clicks":2,"spend":3.82,
             "conversions":0,"conversion_value":0,
             "reason":"Modz is a competing golf cart accessories brand — review only, not recommended as negative"},
            {"category":"CONSIDER","search_term":"icon seat covers",
             "campaigns":"CQC_NBS","impressions":1,"clicks":1,"spend":2.06,
             "conversions":0,"conversion_value":0,
             "reason":"Ambiguous: 'Icon' could refer to Icon golf cart brand (relevant) or a generic seat cover brand (less relevant)"},
        ]
    },
    "GoGather": {
        "total_analyzed": 14,
        "flagged": [
            {"category":"ELIMINATE","search_term":"go cadmium",
             "campaigns":"BFD_GoGather_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":16.13,
             "conversions":0,"conversion_value":0,
             "reason":"Cadmium is event management software, not a corporate planning agency — wrong product category"},
            {"category":"CONSIDER","search_term":"event planner conferences",
             "campaigns":"BFD_GoGather_Google_NB_Search_Pros","impressions":2,"clicks":2,"spend":6.94,
             "conversions":0,"conversion_value":0,
             "reason":"Likely someone seeking a professional development conference FOR event planners, not hiring a planning agency"},
            {"category":"CONSIDER","search_term":"event planning and management",
             "campaigns":"BFD_GoGather_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":6.90,
             "conversions":0,"conversion_value":0,
             "reason":"Very generic — $6.90 spend, zero conversions; broad match may attract researchers rather than B2B buyers"},
            {"category":"COMPETITOR","search_term":"maritz global events",
             "campaigns":"BFD_GoGather_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":6.09,
             "conversions":0,"conversion_value":0,
             "reason":"Maritz Global Events is a competing corporate events/travel management company — review only"},
            {"category":"CONSIDER","search_term":"event planners conference",
             "campaigns":"BFD_GoGather_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":6.38,
             "conversions":0,"conversion_value":0,
             "reason":"Likely someone seeking a conference for event planners (PCMA/MPI-type), not hiring GoGather"},
            {"category":"CONSIDER","search_term":"meeting planners",
             "campaigns":"BFD_GoGather_Google_NB_Search_Pros","impressions":1,"clicks":1,"spend":5.06,
             "conversions":0,"conversion_value":0,
             "reason":"Ambiguous — could be someone who IS a meeting planner (not a buyer) or someone seeking planner services"},
        ]
    },
    "Ripco": {
        "total_analyzed": 6,
        "flagged": []
    },
}

# ── keyword helpers ──────────────────────────────────────────────────────────
def contains_any(term, patterns):
    t = term.lower()
    return any(p.lower() in t for p in patterns)

JOB_TERMS = ["job","jobs","career","careers","hiring","salary","how much do","how to become",
              "certification program","degree program","employment","resume","staffing agency"]
DIY_TERMS  = ["diy","how to make","how to build","homemade","self-made","pattern","tutorial",
              "plans for","build your own","make your own","step by step"]
INFO_TERMS = ["what is ","how does ","why is ","history of ","definition of","explained","guide to ","overview of"]

# ── per-client competitor lists ───────────────────────────────────────────────
ACTIVEWEAR_COMPETITORS = [
    "lululemon","gymshark","under armour","nike","adidas","puma","reebok",
    "vuori","rhone","ten thousand","cuts clothing","represent","stone island",
    "alphalete","gym king","dsquared","acne studios"
]
MORTGAGE_COMPETITORS = [
    "rocket mortgage","quicken loans","better.com","better mortgage","penfed",
    "home xpress","loanify","ps funding","johnson bank","royal united",
    "guidefunding","integrity first","np inc","harbour realty","alphatech",
    "splitero","loandepot","loan depot","newrez","pennymac","chase bank",
    "amerisave","upstart","unlock hea","bd nationwide","prosperity bank",
    "backflip","inland capital","homexpress","wells fargo","chase home",
    "homeexpress","home express mortgage","velocity mortgage","deephaven",
    "angel oak","acra lending","quontic","lima one","kiavi",
]
LEARN_COMPETITORS = [
    # Named competitors found in this week's data
    "brighter strides","triangle aba","hopebridge","caravel","lighthouse autism",
    "lighthouse aba","kennedy aba","proven behavior solutions","bci aba","bci learning",
    "bci autism","journey aba","cabi ","lifespan aba","above and beyond therapy",
    "cortica","autism home support","therapeutic pathways","behavior frontiers",
    "may institute","beacon aba","acorn health","action behavior centers",
    "sunbelt staffing","little star aba","futures west","children making strides",
    "behavioral concepts","cardinal pediatric","mosaic pediatric","hands autism",
    "cbs aba","aligned aba","stepping stones","north carolina behavior",
    "abc center autism","aba centers proven","mosaic autism",
]
ECOPEACO_COMPETITORS = [
    "honest company","the honest company","pipette","huggies natural",
    "esembly","carina organics","kudo","bambi","bambi diapers",
    "pampers pure","seventh generation","andy pandy","eco by naty",
    "dyper","parasol",
]
WE_BUY_COMPETITORS = [
    "ilusso","illusso","exotic car trader","exotic cars trader","j&s autohaus","j and s autohaus",
    "dupont registry","dupontregistry","hgregoire exotic","champagne motor car",
    "select luxury","prestige imports","celebrity motor car","exotic car list",
]

TARGET_CPA = {
    "ASRV":50,"Griffin":50,"Velo":50,"Sunrise":150,"TriActive":150,
    "LEARN":200,"Eurolog":500,"Spotlight":50,"CQC":50,
    "EcoPeaCo CA":50,"EcoPeaCo US":50,"GoGather":50,"RCM":50,"Ripco":50,"We Buy":50
}

# ── threshold filter + dedup ─────────────────────────────────────────────────
def filter_and_dedup(rows):
    agg = defaultdict(lambda: {"impressions":0,"clicks":0,"spend":0.0,"conversions":0.0,"conversion_value":0.0,"campaigns":set()})
    for r in rows:
        st = r.get("search_term")
        if not st:
            continue
        agg[st]["impressions"]      += r.get("impressions", 0)
        agg[st]["clicks"]           += r.get("clicks", 0)
        agg[st]["spend"]            += r.get("spend", 0.0)
        agg[st]["conversions"]      += r.get("conversions", 0.0)
        agg[st]["conversion_value"] += r.get("conversion_value", 0.0)
        camp = r.get("campaign","")
        if camp:
            agg[st]["campaigns"].add(camp)
    filtered = []
    for st, m in agg.items():
        if m["clicks"] >= 1 or m["spend"] >= 1.0:
            filtered.append({
                "search_term": st,
                "impressions": m["impressions"],
                "clicks": m["clicks"],
                "spend": round(m["spend"], 2),
                "conversions": m["conversions"],
                "conversion_value": round(m["conversion_value"], 2),
                "campaigns": ", ".join(sorted(m["campaigns"])),
            })
    return filtered

# ── categorization per client ─────────────────────────────────────────────────
def categorize(r, client_key):
    st      = r["search_term"].lower()
    spend   = r["spend"]
    conv    = r["conversions"]
    camps   = r["campaigns"].lower()
    target_cpa = TARGET_CPA.get(client_key, 50)
    high_spend = target_cpa * 3

    # Rule: never eliminate converting terms
    if conv >= 1 and client_key != "Spotlight":
        # May still be COMPETITOR
        pass

    # ── Griffin Funding ───────────────────────────────────────────────────────
    if client_key == "Griffin":
        # Garbled / wrong DSCR acronyms — searchers confused about the product
        dscr_garbles = ["dsr loan","dsc loan","dcrs loan","drc loan","dsr rates"]
        if any(g in st for g in dscr_garbles) and "dscr" not in st:
            return "ELIMINATE", "Garbled DSCR acronym — searcher likely confused about loan product; wrong-intent traffic"
        if "debt servicing" in st:
            return "ELIMINATE", "Garbled term for debt-service coverage ratio; not a standard mortgage search"
        if "hud 184" in st or "hud184" in st:
            return "ELIMINATE", "HUD 184 is a Native American-specific mortgage program — outside Griffin's non-QM/bank statement target market"
        if contains_any(st, MORTGAGE_COMPETITORS):
            if conv >= 1:
                return "COMPETITOR", "Competitor brand search — converting; do not add as negative without review"
            return "COMPETITOR", "Competitor brand search — review only, not recommended as negative"
        if contains_any(st, JOB_TERMS):
            return "ELIMINATE", "Job-seeking query — no loan purchase intent"
        # Only flag at 3x CPA ($150) to avoid over-flagging competitive mortgage terms
        if conv == 0 and spend > high_spend:
            return "CONSIDER", f"${spend:.2f} spend, zero conversions — high-cost non-converting term; review match type or ad relevance"
        return None, None

    # ── LEARN Behavioral ─────────────────────────────────────────────────────
    if client_key == "LEARN":
        # Competitor ABA providers — review only
        if contains_any(st, LEARN_COMPETITORS):
            return "COMPETITOR", "Competing ABA therapy / behavioral health provider — review only, not recommended as negative"
        # DIAG campaigns: adult self-assessment / autism testing queries
        # These run in campaigns explicitly tagged _DIAG — adults seeking diagnosis for themselves
        # LEARN's target is families of children with autism; adult self-screeners are misaligned
        if "_diag" in camps:
            if conv == 0:
                return "CONSIDER", f"Adult self-diagnostic query in DIAG campaign (${spend:.2f} spend, zero conversions) — verify whether adult self-screeners align with LEARN's child/family enrollment audience"
        if contains_any(st, JOB_TERMS):
            return "ELIMINATE", "Job/career query — no therapy enrollment intent"
        if conv == 0 and spend > high_spend:
            return "CONSIDER", f"Zero conversions with ${spend:.2f} spend — review landing page alignment"
        if conv == 0 and spend > target_cpa:
            return "CONSIDER", f"${spend:.2f} spend, zero conversions — monitor against ${target_cpa} CPA target"
        return None, None

    # ── EcoPeaCo CA & EcoPeaCo US ────────────────────────────────────────────
    if client_key in ("EcoPeaCo CA", "EcoPeaCo US"):
        # Brand exclusions
        if contains_any(st, ["ecopeaco","eco pea co","eco pea","eco peaco","ecopea"]):
            return None, None
        if contains_any(st, ECOPEACO_COMPETITORS):
            if conv >= 1:
                return "COMPETITOR", "Competitor brand — converting; review only, not recommended as negative"
            return "COMPETITOR", "Competitor baby diaper/wipes brand — review only"
        if contains_any(st, JOB_TERMS):
            return "ELIMINATE", "Job-seeking query — no purchase intent"
        if contains_any(st, DIY_TERMS):
            return "ELIMINATE", "DIY/crafting query — not a buyer of finished eco baby products"
        if "cloth diaper service" in st or "diaper service" in st:
            return "ELIMINATE", "Diaper laundry service query — wrong product type; EcoPeaCo sells disposables"
        if "cloth diaper" in st and "organic cloth" not in st:
            return "CONSIDER", "Cloth diaper searcher — different product type from EcoPeaCo's eco-disposables"
        if "training underwear" in st or "training pants" in st or "potty training pants" in st:
            return "CONSIDER", "Training pants query — adjacent but distinct product; verify EcoPeaCo inventory alignment"
        if "wholesale" in st or "bulk" in st:
            return "CONSIDER", "Wholesale/bulk query — may be a retailer buyer rather than DTC parent"
        # Flag generic non-brand terms with spend ≥$4 — in a $50 DTC market, $4 with 0 conv merits review
        if conv == 0 and spend >= 4.0:
            return "CONSIDER", f"${spend:.2f} spend, zero conversions — generic term may attract broad comparison shoppers; review ad relevance"
        return None, None

    # ── GoGather ─────────────────────────────────────────────────────────────
    if client_key == "GoGather":
        if "cadmium" in st or "cad" == st.strip():
            return "ELIMINATE", "Cadmium is event management software, not a planning agency — wrong product category"
        if contains_any(st, ["maritz","anschutz","pcma","mpi meeting","cvent","bizzabo","eventbrite"]):
            return "COMPETITOR", "Competing corporate events / meeting management platform or agency — review only"
        if contains_any(st, JOB_TERMS):
            return "ELIMINATE", "Job-seeking query — not an event planner seeking services"
        if contains_any(st, DIY_TERMS):
            return "ELIMINATE", "DIY event planning query — self-organizing, not seeking professional services"
        if conv == 0 and spend > high_spend:
            return "ELIMINATE", f"Zero conversions with ${spend:.2f} spend — appears off-target"
        if conv == 0 and spend > target_cpa:
            return "CONSIDER", f"${spend:.2f} spend, zero conversions — review keyword relevance"
        return None, None

    # ── RCM (Renewable Credit Management) ────────────────────────────────────
    if client_key == "RCM":
        # Brand exclusion
        if "renewable credit management" in st or st in ("rcm", "rcm tax credits"):
            return None, None
        if "giraffe" in st:
            return "COMPETITOR", "Giraffe Tax Credits is a competing IRA credit brokerage — review only"
        if "residential" in st and ("tax credit" in st or "credit" in st):
            return "ELIMINATE", "Residential housing/tax credit query — wrong product; RCM targets commercial energy credits"
        if "rental" in st and "credit" in st:
            return "ELIMINATE", "Residential rental tax credit — wrong product category for RCM's commercial/energy credit focus"
        if "solar" in st and "residential" in st:
            return "ELIMINATE", "Residential solar query — RCM targets transferable commercial energy credits, not consumer solar"
        if contains_any(st, JOB_TERMS):
            return "ELIMINATE", "Job-seeking query — not a business seeking tax credit services"
        # Flag non-converting terms at $8+ spend — low-volume specialist market where $8 = meaningful signal
        if conv == 0 and spend >= 8.0:
            return "CONSIDER", f"${spend:.2f} spend, zero conversions — review whether search intent aligns with transferable commercial energy credit buyers"
        return None, None

    # ── We Buy Exotics ────────────────────────────────────────────────────────
    if client_key == "We Buy":
        # Brand exclusions
        if "we buy exotics" in st:
            return None, None
        if contains_any(st, WE_BUY_COMPETITORS):
            if conv >= 1:
                return "COMPETITOR", "Competitor exotic car platform — converting; review carefully before adding as negative"
            return "COMPETITOR", "Competitor exotic car buyer/dealer platform — review only"
        if contains_any(st, JOB_TERMS):
            return "ELIMINATE", "Job-seeking query — not a seller looking to sell an exotic vehicle"
        # Clear seller intent — these are the people We Buy Exotics wants; do NOT flag
        seller_signals = ["sell my ","sell a ","sell your ","selling my ","selling a ",
                          "selling an ","selling your ","selling the ","where to sell",
                          "best place to sell","how to sell","who buys","who will buy",
                          "want to sell","looking to sell","need to sell","can you sell",
                          "sell high end","sell luxury","sell exotic","sell super",
                          "sell porsche","sell ferrari","sell lambo","sell bentley",
                          "sell rolls","sell bmw","sell audi","sell mercedes",
                          "i want to sell","sell an "]
        if contains_any(st, seller_signals):
            return None, None
        # Buyer intent (people wanting to PURCHASE a car, not SELL one)
        buyer_signals = ["dealership","dealers near me","dealer near me","buy a ","buy an ",
                         "buy new ","buy used ","for sale near","purchase a ","lease a ",
                         "financing","msrp","invoice price","new 2026","used car lot",
                         "used car dealership","exotic cars near me","near me",
                         "aston martin buy","ferrari buy","lamborghini buy",
                         "buy aston","buy ferrari","buy lambo","buy porsche","buy rolls",
                         "buy bentley","buy mclaren","find aston","aston martin deals",
                         "exotic car deals","car shop","car finder","car agency"]
        if contains_any(st, buyer_signals):
            if conv == 0:
                return "ELIMINATE", "Vehicle purchase/buyer intent — We Buy Exotics BUYS cars; this is the opposite intent"
        if "dealer" in st and conv == 0:
            return "ELIMINATE", "Dealership search — buyer seeking a retailer to purchase from, not a service that buys their car"
        # Valuation informational queries
        if contains_any(st, ["how much is","how much does","value of my","worth of my",
                              "car value","trade in value","trade-in value","what is my car worth",
                              "how much can i get","car worth","kbb","kelley blue book"]):
            if conv == 0:
                return "ELIMINATE", "Car valuation informational query — researching price before deciding whether to sell; no direct transaction signal"
        # Ambiguous high-spend terms (e.g., 'porsche buyers', 'consignment', '911 buyers')
        if conv == 0 and spend > high_spend:
            return "CONSIDER", f"${spend:.2f} spend, zero conversions — ambiguous seller/buyer intent; review ad copy and landing page alignment"
        return None, None

    return None, None

# ── load data from disk ───────────────────────────────────────────────────────
def load_and_process(client_key):
    fname = DISK_FILES.get(client_key)
    if not fname:
        return 0, []
    path = os.path.join(DATA_DIR, fname)
    with open(path) as f:
        data = json.load(f)
    rows = data.get("result", [])
    filtered = filter_and_dedup(rows)
    total_analyzed = len(filtered)

    flagged = []
    for r in filtered:
        cat, reason = categorize(r, client_key)
        if cat:
            # Downgrade ELIMINATE to CONSIDER if the term has conversions
            if cat == "ELIMINATE" and r["conversions"] >= 1:
                cat = "CONSIDER"
                reason = reason + " (has conversions — downgraded from ELIMINATE)"
            flagged.append({**r, "category": cat, "reason": reason})

    return total_analyzed, flagged

# ── load Spotlight from pre-processed JSON ────────────────────────────────────
def load_spotlight():
    with open("/tmp/spotlight_flagged.json") as f:
        data = json.load(f)
    total = data["summary"]["TOTAL_ANALYZED"]
    flagged = data["flagged"]
    # Normalize field names / types
    cleaned = []
    for r in flagged:
        cleaned.append({
            "search_term":    r["search_term"],
            "category":       r["category"],
            "campaigns":      r["campaigns"],
            "impressions":    r["impressions"],
            "clicks":         r["clicks"],
            "spend":          round(float(r["spend"]), 2),
            "conversions":    r["conversions"],
            "conversion_value": round(float(r.get("conversion_value", 0)), 2),
            "reason":         r["reason"],
        })
    return total, cleaned

# ── Excel formatting ──────────────────────────────────────────────────────────
RED_FILL    = PatternFill("solid", fgColor="C00000")
YELLOW_FILL = PatternFill("solid", fgColor="FFFF00")
ORANGE_FILL = PatternFill("solid", fgColor="ED7D31")
HEADER_FILL = PatternFill("solid", fgColor="2E4057")
SUMM_FILL   = PatternFill("solid", fgColor="D9E1F2")

WHITE_BOLD  = Font(bold=True, color="FFFFFF")
BLACK_BOLD  = Font(bold=True, color="000000")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT  = Font(bold=True, size=13)
SUMM_FONT   = Font(bold=True)

def set_col_widths(ws, widths):
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

def write_client_sheet(ws, client_key, flagged_rows, total_analyzed, note=None):
    display = CLIENT_DISPLAY[client_key]
    cat_counts = {"ELIMINATE":0,"CONSIDER":0,"COMPETITOR":0}
    for r in flagged_rows:
        cat_counts[r["category"]] += 1

    # Row 1 — merged title
    ws.merge_cells("A1:J1")
    ws["A1"] = f"{display} — Negative Keyword Report | Last 7 Days: {DATE_RANGE}"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # Rows 3-6 — summary block
    summary = [
        ("Total Search Terms Analyzed", total_analyzed),
        ("Terms to Eliminate",          cat_counts["ELIMINATE"]),
        ("Terms to Consider",           cat_counts["CONSIDER"]),
        ("Competitor Terms",            cat_counts["COMPETITOR"]),
    ]
    for i, (label, val) in enumerate(summary, start=3):
        ws.cell(i, 1, label).font = SUMM_FONT
        ws.cell(i, 2, val).font   = SUMM_FONT
        ws.cell(i, 1).fill = SUMM_FILL
        ws.cell(i, 2).fill = SUMM_FILL

    if note:
        ws.cell(7, 1, note).font = Font(italic=True, color="C00000")

    if not flagged_rows:
        ws.cell(8, 1, "No terms flagged this week.")
        return

    # Row 8 — column headers
    headers = ["Category","Search Term","Matched Campaigns","Impressions","Clicks","Spend","Conversions","Conv. Value","Reason"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(8, col, h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    ws.freeze_panes = "A9"

    # sort: ELIMINATE → CONSIDER → COMPETITOR, then by spend desc
    order = {"ELIMINATE":0,"CONSIDER":1,"COMPETITOR":2}
    flagged_rows.sort(key=lambda x: (order[x["category"]], -x["spend"]))

    for row_idx, r in enumerate(flagged_rows, start=9):
        cat = r["category"]
        ws.cell(row_idx, 1, cat)
        ws.cell(row_idx, 2, r["search_term"])
        ws.cell(row_idx, 3, r["campaigns"])
        ws.cell(row_idx, 4, r["impressions"])
        ws.cell(row_idx, 5, r["clicks"])
        spend_cell = ws.cell(row_idx, 6, r["spend"])
        spend_cell.number_format = '$#,##0.00'
        ws.cell(row_idx, 7, r["conversions"])
        cv_cell = ws.cell(row_idx, 8, r["conversion_value"])
        cv_cell.number_format = '$#,##0.00'
        ws.cell(row_idx, 9, r["reason"])

        cat_cell = ws.cell(row_idx, 1)
        if cat == "ELIMINATE":
            cat_cell.fill = RED_FILL
            cat_cell.font = WHITE_BOLD
        elif cat == "CONSIDER":
            cat_cell.fill = YELLOW_FILL
            cat_cell.font = BLACK_BOLD
        elif cat == "COMPETITOR":
            cat_cell.fill = ORANGE_FILL
            cat_cell.font = WHITE_BOLD
        cat_cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[row_idx].height = 16

    set_col_widths(ws, [14, 36, 40, 12, 8, 10, 13, 13, 60])

# ── main ──────────────────────────────────────────────────────────────────────
def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    summary_stats = {}

    for client_key in TAB_ORDER:
        print(f"Processing {client_key}...", end=" ", flush=True)

        note = None

        if client_key in PRECATEGORIZED:
            pc = PRECATEGORIZED[client_key]
            total_analyzed = pc["total_analyzed"]
            flagged = list(pc["flagged"])  # copy
            note = pc.get("note")

        elif client_key == "Spotlight":
            total_analyzed, flagged = load_spotlight()

        else:
            total_analyzed, flagged = load_and_process(client_key)

        ws = wb.create_sheet(title=CLIENT_DISPLAY[client_key][:31])
        write_client_sheet(ws, client_key, flagged, total_analyzed, note=note)

        cat_counts = {"ELIMINATE":0,"CONSIDER":0,"COMPETITOR":0}
        for fr in flagged:
            cat_counts[fr["category"]] += 1
        summary_stats[client_key] = cat_counts
        print(f"{total_analyzed} terms → E:{cat_counts['ELIMINATE']} C:{cat_counts['CONSIDER']} Comp:{cat_counts['COMPETITOR']}")

    out_path = f"/home/user/routines-ppc/reports/negative-keyword-report-{REPORT_DATE}.xlsx"
    wb.save(out_path)
    print(f"\nSaved: {out_path}")

    print("\n=== SUMMARY FOR SLACK ===")
    for client_key in TAB_ORDER:
        s = summary_stats[client_key]
        print(f"{client_key}: E={s['ELIMINATE']} C={s['CONSIDER']} Comp={s['COMPETITOR']}")

    return out_path, summary_stats

if __name__ == "__main__":
    main()
